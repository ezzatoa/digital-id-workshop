# -*- coding: utf-8 -*-
"""
Full Presentation Builder for:
"الهوية الرقمية الأكاديمية والحضور البحثي العالمي"
Presenter: د. عزت عمر عبدالله أبوعزه
Affiliation: أستاذ مساعد في قسم تقنية الأشعة - كلية العلوم الطبية التطبيقية - جامعة طيبة
Based on the official Taibah University 1448H template.
"""

import sys
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Brand Colors (Taibah University Palette)
NAVY_DEEP = RGBColor(11, 26, 72)       # #0b1a48
NAVY_MEDIUM = RGBColor(18, 42, 107)    # #122a6b
TEAL_EMERALD = RGBColor(0, 168, 135)   # #00a887
CYAN_ACCENT = RGBColor(0, 196, 216)    # #00c4d8
GOLD_ACCENT = RGBColor(217, 119, 6)    # #d97706
RED_ACCENT = RGBColor(220, 38, 38)     # #dc2626
TEXT_DARK = RGBColor(30, 41, 59)       # Slate 800
TEXT_MUTED = RGBColor(100, 116, 139)   # Slate 500
WHITE = RGBColor(255, 255, 255)
CARD_BG = RGBColor(248, 250, 252)      # Slate 50
CARD_BORDER = RGBColor(226, 232, 240)  # Slate 200
HIGHLIGHT_BG = RGBColor(240, 253, 250) # Teal 50
DARK_CARD_BG = RGBColor(18, 32, 85)    # Card on dark bg

BG_TITLE = '_template_extracted/ppt/media/image2.png'
BG_DARK = '_template_extracted/ppt/media/image3.png'
BG_LIGHT = '_template_extracted/ppt/media/image4.png'

prs = pptx.Presentation('خلفية البرامج الاسبوعية 1448.pptx')
slide_width = prs.slide_width
slide_height = prs.slide_height

# Clear existing slides
while len(prs.slides) > 0:
    rId = prs.slides._sldIdLst[0].rId
    prs.part.drop_rel(rId)
    del prs.slides._sldIdLst[0]

blank_layout = prs.slide_layouts[6]

def add_bg(slide, bg_path):
    slide.shapes.add_picture(bg_path, 0, 0, slide_width, slide_height)

def add_header(slide, title_text, category="الهوية الرقمية الأكاديمية والحضور البحثي العالمي"):
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.32), Inches(11.7), Inches(0.35))
    tf_b = tb.text_frame
    tf_b.word_wrap = True
    p_b = tf_b.paragraphs[0]
    p_b.text = f"جامعة طيبة  |  عمادة التطوير والجودة  |  {category}"
    p_b.alignment = PP_ALIGN.RIGHT
    p_b.font.size = Pt(11)
    p_b.font.bold = True
    p_b.font.color.rgb = TEAL_EMERALD
    p_b.font.name = "Arial"

    tx = slide.shapes.add_textbox(Inches(0.8), Inches(0.65), Inches(11.7), Inches(0.75))
    tf = tx.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.alignment = PP_ALIGN.RIGHT
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = NAVY_DEEP
    p.font.name = "Arial"

def add_speaker_notes(slide, notes_text):
    notes_slide = slide.notes_slide
    tf = notes_slide.notes_text_frame
    tf.text = notes_text

def create_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    shape.line.color.rgb = border_color
    shape.line.width = Pt(1.2)
    return shape

def add_card_content(shape, title, items, title_color=NAVY_DEEP, body_color=TEXT_DARK, badge=None):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = Inches(0.25)
    tf.margin_right = Inches(0.25)
    tf.margin_top = Inches(0.22)
    tf.margin_bottom = Inches(0.2)

    p0 = tf.paragraphs[0]
    if badge:
        p0.text = f"[{badge}]  {title}"
    else:
        p0.text = title
    p0.font.size = Pt(15)
    p0.font.bold = True
    p0.font.color.rgb = title_color
    p0.font.name = "Arial"
    p0.alignment = PP_ALIGN.RIGHT
    p0.space_after = Pt(8)

    for item in items:
        p = tf.add_paragraph()
        p.text = f"•  {item}"
        p.font.size = Pt(12)
        p.font.color.rgb = body_color
        p.font.name = "Arial"
        p.alignment = PP_ALIGN.RIGHT
        p.space_after = Pt(5)

def make_section_slide(title, subtitle, module_num, time_range, description, notes):
    slide = prs.slides.add_slide(blank_layout)
    add_bg(slide, BG_DARK)

    pill = create_card(slide, Inches(4.5), Inches(1.8), Inches(4.33), Inches(0.55), bg_color=NAVY_MEDIUM, border_color=CYAN_ACCENT)
    p_pill = pill.text_frame.paragraphs[0]
    p_pill.text = f"المحطة التدريبية {module_num}  |  الزمن المتوقع: {time_range}"
    p_pill.alignment = PP_ALIGN.CENTER
    p_pill.font.size = Pt(13)
    p_pill.font.bold = True
    p_pill.font.color.rgb = CYAN_ACCENT
    p_pill.font.name = "Arial"

    tx = slide.shapes.add_textbox(Inches(1.5), Inches(2.6), Inches(10.33), Inches(1.5))
    tf = tx.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.name = "Arial"

    p_sub = tf.add_paragraph()
    p_sub.text = subtitle
    p_sub.alignment = PP_ALIGN.CENTER
    p_sub.font.size = Pt(18)
    p_sub.font.bold = False
    p_sub.font.color.rgb = CYAN_ACCENT
    p_sub.font.name = "Arial"
    p_sub.space_before = Pt(10)

    box = create_card(slide, Inches(2.2), Inches(4.5), Inches(8.93), Inches(1.5), bg_color=DARK_CARD_BG, border_color=TEAL_EMERALD)
    tf_b = box.text_frame
    tf_b.word_wrap = True
    tf_b.margin_left = Inches(0.4)
    tf_b.margin_right = Inches(0.4)
    tf_b.margin_top = Inches(0.25)
    p_desc = tf_b.paragraphs[0]
    p_desc.text = "الأهداف والمخرجات المباشرة لهذه المحطة:"
    p_desc.font.size = Pt(13)
    p_desc.font.bold = True
    p_desc.font.color.rgb = TEAL_EMERALD
    p_desc.alignment = PP_ALIGN.RIGHT
    p_desc.space_after = Pt(6)

    for d in description:
        p_item = tf_b.add_paragraph()
        p_item.text = f"✔  {d}"
        p_item.font.size = Pt(12)
        p_item.font.color.rgb = WHITE
        p_item.alignment = PP_ALIGN.RIGHT
        p_item.space_after = Pt(4)

    add_speaker_notes(slide, notes)
    return slide

# ==========================================
# SLIDE 1: COVER SLIDE
# ==========================================
s1 = prs.slides.add_slide(blank_layout)
add_bg(s1, BG_TITLE)

# Title box on the right side of cover
tb1 = s1.shapes.add_textbox(Inches(4.5), Inches(2.3), Inches(7.2), Inches(2.3))
tf1 = tb1.text_frame
tf1.word_wrap = True

p_badge = tf1.paragraphs[0]
p_badge.text = "برنامج استقبال وتهيئة أعضاء هيئة التدريس الجدد  |  البرنامج رقم 04"
p_badge.font.size = Pt(13)
p_badge.font.bold = True
p_badge.font.color.rgb = CYAN_ACCENT
p_badge.alignment = PP_ALIGN.RIGHT
p_badge.font.name = "Arial"

p_main = tf1.add_paragraph()
p_main.text = "الهوية الرقمية الأكاديمية\nوالحضور البحثي العالمي"
p_main.font.size = Pt(36)
p_main.font.bold = True
p_main.font.color.rgb = WHITE
p_main.alignment = PP_ALIGN.RIGHT
p_main.font.name = "Arial"
p_main.space_before = Pt(8)

p_sub = tf1.add_paragraph()
p_sub.text = "Academic Digital Identity & Global Research Presence"
p_sub.font.size = Pt(16)
p_sub.font.italic = True
p_sub.font.color.rgb = TEAL_EMERALD
p_sub.alignment = PP_ALIGN.RIGHT
p_sub.font.name = "Arial"
p_sub.space_before = Pt(6)

# Presenter box
pres_box = create_card(s1, Inches(4.5), Inches(4.8), Inches(7.2), Inches(1.5), bg_color=DARK_CARD_BG, border_color=TEAL_EMERALD)
tf_pres = pres_box.text_frame
tf_pres.word_wrap = True
tf_pres.margin_left = Inches(0.3)
tf_pres.margin_right = Inches(0.3)
tf_pres.margin_top = Inches(0.2)

p_pr1 = tf_pres.paragraphs[0]
p_pr1.text = "تقديم: د. عزت عمر عبدالله أبوعزه"
p_pr1.font.size = Pt(16)
p_pr1.font.bold = True
p_pr1.font.color.rgb = WHITE
p_pr1.alignment = PP_ALIGN.RIGHT
p_pr1.font.name = "Arial"

p_pr2 = tf_pres.add_paragraph()
p_pr2.text = "أستاذ مساعد في قسم تقنية الأشعة - كلية العلوم الطبية التطبيقية"
p_pr2.font.size = Pt(13)
p_pr2.font.color.rgb = CYAN_ACCENT
p_pr2.alignment = PP_ALIGN.RIGHT
p_pr2.font.name = "Arial"
p_pr2.space_before = Pt(3)

p_pr3 = tf_pres.add_paragraph()
p_pr3.text = "جامعة طيبة  |  الفصل الدراسي الثاني 1448هـ - 2026م"
p_pr3.font.size = Pt(11)
p_pr3.font.color.rgb = TEXT_MUTED
p_pr3.alignment = PP_ALIGN.RIGHT
p_pr3.font.name = "Arial"
p_pr3.space_before = Pt(3)

add_speaker_notes(s1, """[التوقيت المقترح: 00:00 - 00:03 | الدقائق الأولى]
• الترحيب بالحضور الكريم من الزملاء والزميلات أعضاء هيئة التدريس الجدد بجامعة طيبة.
• التعريف بنفسي: د. عزت عمر عبدالله أبوعزه، أستاذ مساعد بقسم تقنية الأشعة بكلية العلوم الطبية التطبيقية.
• التمهيد لموضوع الورشة: الهوية الرقمية الأكاديمية لم تعد مجرد ترف أو سيرة ذاتية إلكترونية، بل أصبحت البنية التحتية الأساسية لحضور الباحث عالمياً وركيزة رئيسية في سمعة وتصنيف جامعة طيبة في المحافل الدولية.
• تنويه: هذه الورشة تفاعلية 100%، وسنعتمد على منصة ويب تطبيقية خاصة بالورشة تضم محاكيات رقمية ومختبر ذكاء اصطناعي منضبط لتمكين الجميع من التطبيق المباشر.""")

print("Slide 1 generated.")
