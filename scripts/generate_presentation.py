# -*- coding: utf-8 -*-
"""
Generate a complete, professional PowerPoint presentation for the workshop:
"الهوية الرقمية الأكاديمية والحضور البحثي العالمي"
Presenter: د. عزت عمر عبدالله أبوعزه
Affiliation: أستاذ مساعد في قسم تقنية الأشعة - كلية العلوم الطبية التطبيقية - جامعة طيبة
Based on the official Taibah University 1448H template.
"""

import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Brand Colors (Taibah University Palette)
NAVY_DEEP = RGBColor(11, 26, 72)      # #0b1a48
NAVY_MEDIUM = RGBColor(18, 42, 107)   # #122a6b
TEAL_EMERALD = RGBColor(0, 168, 135)  # #00a887
CYAN_ACCENT = RGBColor(0, 196, 216)   # #00c4d8
GOLD_ACCENT = RGBColor(234, 153, 33)  # #ea9921
TEXT_DARK = RGBColor(30, 41, 59)      # Slate 800
TEXT_MUTED = RGBColor(100, 116, 139)  # Slate 500
WHITE = RGBColor(255, 255, 255)
CARD_BG = RGBColor(248, 250, 252)     # Slate 50
CARD_BORDER = RGBColor(226, 232, 240) # Slate 200
HIGHLIGHT_BG = RGBColor(240, 253, 250) # Teal 50

# Image Background Paths
BG_TITLE = '_template_extracted/ppt/media/image2.png'
BG_DARK = '_template_extracted/ppt/media/image3.png'
BG_LIGHT = '_template_extracted/ppt/media/image4.png'

prs = pptx.Presentation('خلفية البرامج الاسبوعية 1448.pptx')
slide_width = prs.slide_width
slide_height = prs.slide_height

# Remove any empty/placeholder slides from template
while len(prs.slides) > 0:
    rId = prs.slides._sldIdLst[0].rId
    prs.part.drop_rel(rId)
    del prs.slides._sldIdLst[0]

blank_layout = prs.slide_layouts[6]

def add_bg(slide, bg_path):
    slide.shapes.add_picture(bg_path, 0, 0, slide_width, slide_height)

def add_header(slide, title_text, category="الهوية الرقمية الأكاديمية والحضور البحثي العالمي"):
    # Header container on light background
    # Top bar badge
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
    tf_b = tb.text_frame
    tf_b.word_wrap = True
    p_b = tf_b.paragraphs[0]
    p_b.text = f"جامعة طيبة  |  عمادة التطوير والجودة  |  {category}"
    p_b.alignment = PP_ALIGN.RIGHT
    p_b.font.size = Pt(11)
    p_b.font.bold = True
    p_b.font.color.rgb = TEAL_EMERALD
    p_b.font.name = "Arial"

    # Slide Main Title
    tx = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.8))
    tf = tx.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.alignment = PP_ALIGN.RIGHT
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = NAVY_DEEP
    p.font.name = "Arial"

def add_speaker_notes(slide, notes_text):
    notes_slide = slide.notes_slide
    tf = notes_slide.notes_text_frame
    tf.text = notes_text

def create_card(slide, left, top, width, height, bg_color=WHITE, border_color=CARD_BORDER):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    shape.line.color.rgb = border_color
    shape.line.width = Pt(1.2)
    return shape

print("Helper functions ready.")
